from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook

from app.profiling.excel_profiler import ExcelProfiler


def _create_workbook(path: Path) -> None:
    workbook = Workbook()

    customers = workbook.active
    customers.title = "Customers"

    customers.append(
        [
            "customer_id",
            "status",
            "balance",
            "signup_date",
            "last_seen",
            "nickname",
        ]
    )

    customers.append(
        [
            1001,
            "ACTIVE",
            10.50,
            date(2026, 1, 10),
            datetime(2026, 1, 10, 8, 30),
            "Al",
        ]
    )

    customers.append(
        [
            1002,
            "INACTIVE",
            25.75,
            date(2026, 2, 15),
            datetime(2026, 2, 15, 14, 45),
            None,
        ]
    )

    customers.append(
        [
            1003,
            "ACTIVE",
            5.25,
            date(2026, 3, 20),
            datetime(2026, 3, 20, 9, 15),
            "  Bea  ",
        ]
    )

    workbook.create_sheet("Orders")

    workbook.save(path)
    workbook.close()


def test_excel_profiler_profiles_worksheet(
    tmp_path: Path,
):
    source = tmp_path / "customers.xlsx"
    _create_workbook(source)

    profiler = ExcelProfiler()

    profiles = profiler.profile(
        source,
        worksheet_name="Customers",
        field_types={
            "customer_id": "integer",
            "status": "string",
            "balance": "decimal",
            "signup_date": "date",
            "last_seen": "datetime",
            "nickname": "string",
        },
    )

    assert len(profiles) == 6

    by_name = {
        profile.field_name: profile
        for profile in profiles
    }

    customer_id = by_name["customer_id"]

    assert customer_id.row_count == 3
    assert customer_id.null_count == 0
    assert (
        customer_id.null_percentage
        == Decimal("0.0000")
    )
    assert customer_id.distinct_count == 3
    assert (
        customer_id.distinct_percentage
        == Decimal("100.0000")
    )
    assert customer_id.minimum_value == "1001"
    assert customer_id.maximum_value == "1003"

    status = by_name["status"]

    assert status.distinct_count == 2
    assert (
        status.distinct_percentage
        == Decimal("66.6667")
    )
    assert status.minimum_length == 6
    assert status.maximum_length == 8

    nickname = by_name["nickname"]

    assert nickname.null_count == 1
    assert (
        nickname.null_percentage
        == Decimal("33.3333")
    )
    assert nickname.minimum_length == 2
    assert nickname.maximum_length == 3

    balance = by_name["balance"]

    assert balance.minimum_value == "5.25"
    assert balance.maximum_value == "25.75"

    signup_date = by_name["signup_date"]

    assert signup_date.minimum_value == "2026-01-10"
    assert signup_date.maximum_value == "2026-03-20"

    last_seen = by_name["last_seen"]

    assert (
        last_seen.minimum_value
        == "2026-01-10T08:30:00"
    )
    assert (
        last_seen.maximum_value
        == "2026-03-20T09:15:00"
    )


def test_excel_profiler_profiles_only_requested_worksheet(
    tmp_path: Path,
):
    source = tmp_path / "customers.xlsx"
    _create_workbook(source)

    profiler = ExcelProfiler()

    profiles = profiler.profile(
        source,
        worksheet_name="Customers",
    )

    field_names = {
        profile.field_name
        for profile in profiles
    }

    assert "customer_id" in field_names
    assert len(profiles) == 6


def test_excel_profiler_rejects_missing_worksheet(
    tmp_path: Path,
):
    source = tmp_path / "customers.xlsx"
    _create_workbook(source)

    profiler = ExcelProfiler()

    try:
        profiler.profile(
            source,
            worksheet_name="Missing",
        )
    except ValueError as exc:
        assert str(exc) == (
            "Worksheet not found in Excel workbook."
        )
    else:
        raise AssertionError(
            "Expected ValueError."
        )


def test_excel_profiler_rejects_non_xlsx_file(
    tmp_path: Path,
):
    source = tmp_path / "customers.csv"
    source.write_text(
        "customer_id\n1001\n",
        encoding="utf-8",
    )

    profiler = ExcelProfiler()

    try:
        profiler.profile(
            source,
            worksheet_name="Customers",
        )
    except ValueError as exc:
        assert str(exc) == (
            "Excel profiler only accepts .xlsx files."
        )
    else:
        raise AssertionError(
            "Expected ValueError."
        )


def test_excel_profiler_rejects_blank_header(
    tmp_path: Path,
):
    source = tmp_path / "blank_header.xlsx"

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Customers"
    worksheet.append(
        [
            "customer_id",
            None,
        ]
    )
    worksheet.append(
        [
            1001,
            "ACTIVE",
        ]
    )
    workbook.save(source)
    workbook.close()

    profiler = ExcelProfiler()

    try:
        profiler.profile(
            source,
            worksheet_name="Customers",
        )
    except ValueError as exc:
        assert (
            "blank column names"
            in str(exc)
        )
    else:
        raise AssertionError(
            "Expected ValueError."
        )


def test_excel_profiler_rejects_duplicate_header(
    tmp_path: Path,
):
    source = tmp_path / "duplicate_header.xlsx"

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Customers"
    worksheet.append(
        [
            "customer_id",
            "customer_id",
        ]
    )
    worksheet.append(
        [
            1001,
            1002,
        ]
    )
    workbook.save(source)
    workbook.close()

    profiler = ExcelProfiler()

    try:
        profiler.profile(
            source,
            worksheet_name="Customers",
        )
    except ValueError as exc:
        assert (
            "duplicate column names"
            in str(exc)
        )
    else:
        raise AssertionError(
            "Expected ValueError."
        )