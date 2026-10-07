"""HTTP tests for dictionary request validation and project filtering."""

import csv
import io
import uuid
from unittest.mock import MagicMock

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.api.dependencies.auth import (
    AuthenticatedSession,
    require_authenticated_session,
)
from app.db.session import get_db
from app.main import app
from app.models.user import User


client = TestClient(app)


@pytest.fixture(autouse=True)
def authenticated_dictionary_requests():
    """Authenticate requests in this test module without creating a real user."""
    test_user = User(
        id=uuid.uuid4(),
        username="dictionary_test_user",
        password_hash="test-only-placeholder",
        is_active=True,
    )

    def override_authentication():
        return AuthenticatedSession(
            user=test_user,
            token="dictionary-test-only-token",
        )

    app.dependency_overrides[require_authenticated_session] = (
        override_authentication
    )

    try:
        yield
    finally:
        app.dependency_overrides.pop(
            require_authenticated_session,
            None,
        )


def test_dictionary_requires_project_id():
    response = client.get("/api/v1/dictionary")

    assert response.status_code == 422


def test_dictionary_rejects_invalid_project_id():
    response = client.get(
        "/api/v1/dictionary",
        params={"project_id": "not-a-uuid"},
    )

    assert response.status_code == 422


def test_dictionary_filters_by_project():
    project_id = uuid.uuid4()
    db = MagicMock()

    # Return no rows; this test checks the generated SQL filter.
    db.execute.return_value.mappings.return_value.all.return_value = []

    def override_get_db():
        yield db

    app.dependency_overrides[get_db] = override_get_db

    try:
        response = client.get(
            "/api/v1/dictionary",
            params={"project_id": str(project_id)},
        )

        assert response.status_code == 200
        assert response.json() == []

        statement = db.execute.call_args.args[0]
        sql = str(statement.compile())

        assert "data_sources.project_id" in sql
        assert project_id in statement.compile().params.values()
    finally:
        app.dependency_overrides.pop(get_db, None)

def test_dictionary_filters_non_cde_fields():
    project_id = uuid.uuid4()
    db = MagicMock()
    db.execute.return_value.mappings.return_value.all.return_value = []

    def override_get_db():
        yield db

    app.dependency_overrides[get_db] = override_get_db

    try:
        response = client.get(
            "/api/v1/dictionary",
            params={
                "project_id": str(project_id),
                "is_cde": "false",
            },
        )

        assert response.status_code == 200
        assert response.json() == []

        statement = db.execute.call_args.args[0]
        compiled = statement.compile()
        sql = str(compiled)

        assert "coalesce" in sql.lower()
        assert "field_governance_metadata.is_cde" in sql
        assert False in compiled.params.values()
    finally:
        app.dependency_overrides.pop(get_db, None)


def test_dictionary_combines_source_department_and_owner_filters():
    project_id = uuid.uuid4()
    source_id = uuid.uuid4()
    db = MagicMock()
    db.execute.return_value.mappings.return_value.all.return_value = []

    def override_get_db():
        yield db

    app.dependency_overrides[get_db] = override_get_db

    try:
        response = client.get(
            "/api/v1/dictionary",
            params={
                "project_id": str(project_id),
                "data_source_id": str(source_id),
                "department": "Operations",
                "data_owner": "Test Owner",
            },
        )

        assert response.status_code == 200
        assert response.json() == []

        statement = db.execute.call_args.args[0]
        compiled = statement.compile()
        sql = str(compiled)
        values = compiled.params.values()

        assert "data_sources.project_id" in sql
        assert "data_sources.id" in sql
        assert "field_governance_metadata.department" in sql
        assert "field_governance_metadata.data_owner" in sql
        assert project_id in values
        assert source_id in values
        assert "Operations" in values
        assert "Test Owner" in values
    finally:
        app.dependency_overrides.pop(get_db, None)
def test_dictionary_returns_discovered_technical_metadata():
    project_id = uuid.uuid4()
    field_id = uuid.uuid4()
    source_object_id = uuid.uuid4()
    data_source_id = uuid.uuid4()
    db = MagicMock()

    row = {
        "field_id": field_id,
        "field_name": "customer_code",
        "ordinal_position": 2,
        "native_data_type": "character varying(50)",
        "normalized_data_type": "string",
        "max_length": 50,
        "numeric_precision": None,
        "numeric_scale": None,
        "default_value": "'UNKNOWN'",
        "source_comment": "External customer code",
        "is_nullable": False,
        "is_primary_key": False,
        "is_unique": True,
        "source_object_id": source_object_id,
        "object_name": "customers",
        "object_type": "table",
        "catalog_name": None,
        "schema_name": "public",
        "native_name": "public.customers",
        "object_description": "Customer master data",
        "object_row_count": 1250,
        "data_source_id": data_source_id,
        "data_source_name": "Test PostgreSQL",
        "source_type": "postgresql",
        "department": None,
        "data_owner": None,
        "is_cde": False,
    }

    db.execute.return_value.mappings.return_value.all.return_value = [row]

    def override_get_db():
        yield db

    app.dependency_overrides[get_db] = override_get_db

    try:
        response = client.get(
            "/api/v1/dictionary",
            params={"project_id": str(project_id)},
        )

        assert response.status_code == 200
        field = response.json()[0]

        assert field["max_length"] == 50
        assert field["numeric_precision"] is None
        assert field["numeric_scale"] is None
        assert field["default_value"] == "'UNKNOWN'"
        assert field["source_comment"] == "External customer code"
        assert field["catalog_name"] is None
        assert field["schema_name"] == "public"
        assert field["native_name"] == "public.customers"
        assert field["object_description"] == "Customer master data"
        assert field["object_row_count"] == 1250
    finally:
        app.dependency_overrides.pop(get_db, None)

def test_dictionary_export_requires_project_id_and_format():
    response = client.get("/api/v1/dictionary/export")

    assert response.status_code == 422


def test_dictionary_export_rejects_invalid_format():
    response = client.get(
        "/api/v1/dictionary/export",
        params={
            "project_id": str(uuid.uuid4()),
            "format": "pdf",
        },
    )

    assert response.status_code == 422

def test_dictionary_export_csv_returns_metadata_and_protects_formulas():
    project_id = uuid.uuid4()
    db = MagicMock()

    # Project access is granted when require_project_access() finds a row.
    db.scalar.return_value = uuid.uuid4()

    row = {
        "data_source_name": "Customer PostgreSQL",
        "source_type": "postgresql",
        "catalog_name": None,
        "schema_name": "public",
        "object_name": "customers",
        "object_type": "table",
        "native_name": "public.customers",
        "object_description": "Customer master data",
        "object_row_count": 1250,
        "field_name": "customer_code",
        "ordinal_position": 1,
        "native_data_type": "character varying",
        "normalized_data_type": "string",
        "max_length": 50,
        "numeric_precision": None,
        "numeric_scale": None,
        "is_nullable": False,
        "is_primary_key": True,
        "is_unique": True,
        "default_value": "'UNKNOWN'",
        "source_comment": "External customer code",
        "business_name": "Customer Code",
        "business_definition": "=HYPERLINK(\"https://example.invalid\")",
        "department": "Operations",
        "data_owner": "Data Owner",
        "data_steward": "Data Steward",
        "business_process": "Customer Management",
        "system_of_record": "CRM",
        "is_cde": True,
        "classification": "Internal",
        "approval_status": "draft",
        "notes": "Export test",
    }

    db.execute.return_value.mappings.return_value.all.return_value = [row]

    def override_get_db():
        yield db

    app.dependency_overrides[get_db] = override_get_db

    try:
        response = client.get(
            "/api/v1/dictionary/export",
            params={
                "project_id": str(project_id),
                "format": "csv",
            },
        )

        assert response.status_code == 200
        assert (
            f'filename="geovaris-data-dictionary-{project_id}.csv"'
            in response.headers["content-disposition"]
        )

        reader = csv.reader(
            io.StringIO(response.content.decode("utf-8-sig"))
        )
        rows = list(reader)

        assert rows[0][0] == "Data Source"
        assert rows[0][-1] == "Notes"
        assert rows[1][0] == "Customer PostgreSQL"
        assert rows[0][2] == "Catalog"
        assert rows[1][2] == ""
        assert rows[1][9] == "customer_code"
        assert rows[1][22] == "'=HYPERLINK(\"https://example.invalid\")"
        assert rows[1][28] == "True"
        assert rows[1][30] == "draft"
    finally:
        app.dependency_overrides.pop(get_db, None)

def test_dictionary_export_xlsx_returns_valid_workbook():
    project_id = uuid.uuid4()
    db = MagicMock()

    db.scalar.return_value = uuid.uuid4()

    row = {
        "data_source_name": "Customer PostgreSQL",
        "source_type": "postgresql",
        "catalog_name": None,
        "schema_name": "public",
        "object_name": "customers",
        "object_type": "table",
        "native_name": "public.customers",
        "object_description": "Customer master data",
        "object_row_count": 1250,
        "field_name": "customer_code",
        "ordinal_position": 1,
        "native_data_type": "character varying",
        "normalized_data_type": "string",
        "max_length": 50,
        "numeric_precision": None,
        "numeric_scale": None,
        "is_nullable": False,
        "is_primary_key": True,
        "is_unique": True,
        "default_value": "'UNKNOWN'",
        "source_comment": "External customer code",
        "business_name": "Customer Code",
        "business_definition": "=1+1",
        "department": "Operations",
        "data_owner": "Data Owner",
        "data_steward": "Data Steward",
        "business_process": "Customer Management",
        "system_of_record": "CRM",
        "is_cde": True,
        "classification": "Internal",
        "approval_status": "draft",
        "notes": "Excel export test",
    }

    db.execute.return_value.mappings.return_value.all.return_value = [row]

    def override_get_db():
        yield db

    app.dependency_overrides[get_db] = override_get_db

    try:
        response = client.get(
            "/api/v1/dictionary/export",
            params={
                "project_id": str(project_id),
                "format": "xlsx",
            },
        )

        assert response.status_code == 200
        assert (
            f'filename="geovaris-data-dictionary-{project_id}.xlsx"'
            in response.headers["content-disposition"]
        )

        workbook = load_workbook(
            io.BytesIO(response.content),
            read_only=True,
            data_only=False,
        )

        assert workbook.sheetnames == ["Data Dictionary"]

        worksheet = workbook["Data Dictionary"]
        rows = list(worksheet.iter_rows(values_only=True))

        assert rows[0][0] == "Data Source"
        assert rows[0][-1] == "Notes"
        assert rows[1][0] == "Customer PostgreSQL"
        assert rows[0][2] == "Catalog"
        assert rows[1][2] is None
        assert rows[1][9] == "customer_code"
        assert rows[1][22] == "'=1+1"
        assert rows[1][28] is True
        assert rows[1][30] == "draft"
    finally:
        app.dependency_overrides.pop(get_db, None)

def test_dictionary_export_denies_unauthorized_project():
    project_id = uuid.uuid4()
    db = MagicMock()

    # No project access row is found for the authenticated user.
    db.scalar.return_value = None

    def override_get_db():
        yield db

    app.dependency_overrides[get_db] = override_get_db

    try:
        response = client.get(
            "/api/v1/dictionary/export",
            params={
                "project_id": str(project_id),
                "format": "csv",
            },
        )

        assert response.status_code == 404
        assert response.json()["detail"] == "Project not found."
        db.execute.assert_not_called()
    finally:
        app.dependency_overrides.pop(get_db, None)
