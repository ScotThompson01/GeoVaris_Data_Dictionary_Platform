from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook


@dataclass(frozen=True)
class FieldProfile:
    field_name: str
    row_count: int
    null_count: int
    null_percentage: Decimal
    distinct_count: int
    distinct_percentage: Decimal
    minimum_value: str | None
    maximum_value: str | None
    minimum_length: int | None
    maximum_length: int | None


class ExcelProfiler:
    def profile(
        self,
        source: Path,
        worksheet_name: str,
        field_types: dict[str, str | None] | None = None,
    ) -> list[FieldProfile]:
        """
        Profile one worksheet from an Excel workbook.

        The workbook is opened read-only. Raw worksheet rows are used
        only to calculate aggregate statistics and are not persisted.
        """

        self._validate(source)

        field_types = field_types or {}

        workbook = load_workbook(
            filename=source,
            read_only=True,
            data_only=True,
        )

        try:
            if worksheet_name not in workbook.sheetnames:
                raise ValueError(
                    "Worksheet not found in Excel workbook."
                )

            worksheet = workbook[worksheet_name]

            rows = worksheet.iter_rows(
                values_only=True
            )

            try:
                header_row = next(rows)
            except StopIteration:
                raise ValueError(
                    "Excel worksheet does not contain a header."
                )

            field_names = [
                self._header_name(value)
                for value in header_row
            ]

            self._validate_headers(field_names)

            values_by_field: dict[
                str,
                list[object | None],
            ] = {
                field_name: []
                for field_name in field_names
            }

            row_count = 0

            for row in rows:
                row_count += 1

                for index, field_name in enumerate(
                    field_names
                ):
                    value = (
                        row[index]
                        if index < len(row)
                        else None
                    )

                    values_by_field[field_name].append(
                        self._normalize_value(value)
                    )

            profiles: list[FieldProfile] = []

            for field_name in field_names:
                values = values_by_field[field_name]

                populated_values = [
                    value
                    for value in values
                    if value is not None
                ]

                null_count = (
                    row_count
                    - len(populated_values)
                )

                null_percentage = self._percentage(
                    numerator=null_count,
                    denominator=row_count,
                )

                distinct_count = len(
                    {
                        self._string_value(value)
                        for value in populated_values
                    }
                )

                distinct_percentage = self._percentage(
                    numerator=distinct_count,
                    denominator=row_count,
                )

                minimum_value, maximum_value = (
                    self._typed_min_max(
                        values=populated_values,
                        normalized_data_type=field_types.get(
                            field_name
                        ),
                    )
                )

                lengths = [
                    len(self._string_value(value))
                    for value in populated_values
                ]

                minimum_length = (
                    min(lengths)
                    if lengths
                    else None
                )

                maximum_length = (
                    max(lengths)
                    if lengths
                    else None
                )

                profiles.append(
                    FieldProfile(
                        field_name=field_name,
                        row_count=row_count,
                        null_count=null_count,
                        null_percentage=null_percentage,
                        distinct_count=distinct_count,
                        distinct_percentage=distinct_percentage,
                        minimum_value=minimum_value,
                        maximum_value=maximum_value,
                        minimum_length=minimum_length,
                        maximum_length=maximum_length,
                    )
                )

            return profiles

        finally:
            workbook.close()

    @staticmethod
    def _validate(source: Path) -> None:
        if not source.exists():
            raise FileNotFoundError(
                f"Excel file does not exist: {source}"
            )

        if not source.is_file():
            raise ValueError(
                "Excel source must be a file."
            )

        if source.suffix.lower() != ".xlsx":
            raise ValueError(
                "Excel profiler only accepts .xlsx files."
            )

    @staticmethod
    def _header_name(
        value: object | None,
    ) -> str:
        if value is None:
            return ""

        return str(value).strip()

    @staticmethod
    def _validate_headers(
        field_names: list[str],
    ) -> None:
        if not field_names:
            raise ValueError(
                "Excel worksheet does not contain a header."
            )

        if any(
            not field_name
            for field_name in field_names
        ):
            raise ValueError(
                "Excel worksheet contains one or more "
                "blank column names."
            )

        if len(field_names) != len(set(field_names)):
            raise ValueError(
                "Excel worksheet contains duplicate "
                "column names."
            )

    @staticmethod
    def _normalize_value(
        value: object | None,
    ) -> object | None:
        if value is None:
            return None

        if isinstance(value, str):
            stripped = value.strip()

            return (
                stripped
                if stripped != ""
                else None
            )

        return value

    @staticmethod
    def _string_value(
        value: object,
    ) -> str:
        if isinstance(value, datetime):
            return value.isoformat()

        if isinstance(value, date):
            return value.isoformat()

        if isinstance(value, bool):
            return str(value).lower()

        return str(value)

    @staticmethod
    def _percentage(
        numerator: int,
        denominator: int,
    ) -> Decimal:
        if denominator == 0:
            return Decimal("0.0000")

        percentage = (
            Decimal(numerator)
            / Decimal(denominator)
            * Decimal("100")
        )

        return percentage.quantize(
            Decimal("0.0001")
        )

    @classmethod
    def _typed_min_max(
        cls,
        values: list[object],
        normalized_data_type: str | None,
    ) -> tuple[str | None, str | None]:
        if not values:
            return None, None

        try:
            if normalized_data_type == "integer":
                typed_values = [
                    int(value)
                    for value in values
                ]

                return (
                    str(min(typed_values)),
                    str(max(typed_values)),
                )

            if normalized_data_type == "decimal":
                typed_values = [
                    Decimal(str(value))
                    for value in values
                ]

                return (
                    str(min(typed_values)),
                    str(max(typed_values)),
                )

            if normalized_data_type == "date":
                typed_values = [
                    cls._as_date(value)
                    for value in values
                ]

                return (
                    min(typed_values).isoformat(),
                    max(typed_values).isoformat(),
                )

            if normalized_data_type == "datetime":
                typed_values = [
                    cls._as_datetime(value)
                    for value in values
                ]

                return (
                    min(typed_values).isoformat(),
                    max(typed_values).isoformat(),
                )

        except (
            ValueError,
            TypeError,
            InvalidOperation,
        ):
            pass

        string_values = [
            cls._string_value(value)
            for value in values
        ]

        return (
            min(string_values),
            max(string_values),
        )

    @staticmethod
    def _as_date(
        value: object,
    ) -> date:
        if isinstance(value, datetime):
            return value.date()

        if isinstance(value, date):
            return value

        return datetime.strptime(
            str(value),
            "%Y-%m-%d",
        ).date()

    @staticmethod
    def _as_datetime(
        value: object,
    ) -> datetime:
        if isinstance(value, datetime):
            return value

        if isinstance(value, date):
            return datetime.combine(
                value,
                datetime.min.time(),
            )

        return datetime.fromisoformat(
            str(value)
        )