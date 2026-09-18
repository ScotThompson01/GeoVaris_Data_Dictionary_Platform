from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

from app.connectors.base import (
    DiscoveredField,
    DiscoveredObject,
)


class ExcelConnector:
    """
    Read-only metadata discovery for Excel .xlsx workbooks.

    Each worksheet is represented as a separate DiscoveredObject.
    Workbook contents are inspected in memory and are never modified.
    """

    connector_name = "excel"
    connector_version = "0.2.0"

    def validate(
        self,
        source: Path,
    ) -> None:
        if not source.exists():
            raise FileNotFoundError(
                f"Excel file does not exist: {source}"
            )

        if not source.is_file():
            raise ValueError(
                "Excel source must be a file."
            )

        if source.suffix.lower() != ".xlsx":
            raise ValueError(
                "Excel connector only accepts .xlsx files."
            )

    def discover_workbook(
        self,
        source: Path,
    ) -> list[DiscoveredObject]:
        self.validate(source)

        workbook = load_workbook(
            filename=source,
            read_only=True,
            data_only=True,
        )

        try:
            discovered_objects = [
                self._discover_worksheet(
                    workbook_name=source.name,
                    worksheet=worksheet,
                )
                for worksheet in workbook.worksheets
            ]
        finally:
            workbook.close()

        if not discovered_objects:
            raise ValueError(
                "Excel workbook does not contain any worksheets."
            )

        return discovered_objects

    def _discover_worksheet(
        self,
        workbook_name: str,
        worksheet,
    ) -> DiscoveredObject:
        rows = worksheet.iter_rows(
            values_only=True,
        )

        try:
            header_row = next(rows)
        except StopIteration as exc:
            raise ValueError(
                f"Worksheet '{worksheet.title}' is empty."
            ) from exc

        field_names = self._normalize_headers(
            worksheet.title,
            header_row,
        )

        values_by_field: list[list[object]] = [
            []
            for _ in field_names
        ]

        row_count = 0

        for row in rows:
            row_count += 1

            for index in range(
                len(field_names)
            ):
                value = (
                    row[index]
                    if index < len(row)
                    else None
                )

                if value is not None:
                    values_by_field[
                        index
                    ].append(value)

        fields = []

        for index, field_name in enumerate(
            field_names,
            start=1,
        ):
            values = values_by_field[
                index - 1
            ]

            inferred_type = self._infer_type(
                values
            )

            fields.append(
                DiscoveredField(
                    field_name=field_name,
                    ordinal_position=index,
                    native_data_type=inferred_type,
                    normalized_data_type=inferred_type,
                    is_nullable=(
                        len(values) < row_count
                    ),
                )
            )

        return DiscoveredObject(
            object_type="worksheet",
            object_name=worksheet.title,
            native_name=(
                f"{workbook_name}:"
                f"{worksheet.title}"
            ),
            row_count=row_count,
            fields=fields,
        )

    def _normalize_headers(
        self,
        worksheet_name: str,
        header_row: tuple[object, ...],
    ) -> list[str]:
        field_names = [
            (
                str(value).strip()
                if value is not None
                else ""
            )
            for value in header_row
        ]

        if not field_names or all(
            not field_name
            for field_name in field_names
        ):
            raise ValueError(
                f"Worksheet '{worksheet_name}' "
                "does not contain a header."
            )

        if any(
            not field_name
            for field_name in field_names
        ):
            raise ValueError(
                f"Worksheet '{worksheet_name}' "
                "contains one or more blank "
                "column names."
            )

        if len(field_names) != len(
            set(field_names)
        ):
            raise ValueError(
                f"Worksheet '{worksheet_name}' "
                "contains duplicate column names."
            )

        return field_names

    def _infer_type(
        self,
        values: list[object],
    ) -> str:
        if not values:
            return "string"

        if all(
            isinstance(value, bool)
            for value in values
        ):
            return "boolean"

        if all(
            isinstance(value, int)
            and not isinstance(value, bool)
            for value in values
        ):
            return "integer"

        if all(
            isinstance(
                value,
                (int, float, Decimal),
            )
            and not isinstance(value, bool)
            for value in values
        ):
            return "decimal"

        if all(
            isinstance(value, datetime)
            for value in values
        ):
            if all(
                value.hour == 0
                and value.minute == 0
                and value.second == 0
                and value.microsecond == 0
                for value in values
            ):
                return "date"

            return "datetime"

        if all(
            isinstance(value, date)
            and not isinstance(
                value,
                datetime,
            )
            for value in values
        ):
            return "date"

        return "string"